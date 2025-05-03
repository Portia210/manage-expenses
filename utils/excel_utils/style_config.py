"""Excel styling configuration."""
from dataclasses import dataclass
import openpyxl.styles

@dataclass
class ExcelColors:
    """Excel color configurations."""
    HEADER_COLOR: str = '0070e0'
    FIRST_COLUMN_COLOR: str = 'E6E6E6'

@dataclass
class ColumnConfig:
    """Column width configuration."""
    DEFAULT_PADDING: int = 4
    MAX_WIDTH: int = 40

class StyleFactory:
    """Factory for creating Excel styles."""
    
    @staticmethod
    def create_header_fill() -> openpyxl.styles.PatternFill:
        """Create header fill style."""
        return openpyxl.styles.PatternFill(
            start_color=ExcelColors.HEADER_COLOR,
            end_color=ExcelColors.HEADER_COLOR,
            fill_type='solid'
        )
    
    @staticmethod
    def create_first_column_fill() -> openpyxl.styles.PatternFill:
        """Create first column fill style."""
        return openpyxl.styles.PatternFill(
            start_color=ExcelColors.FIRST_COLUMN_COLOR,
            end_color=ExcelColors.FIRST_COLUMN_COLOR,
            fill_type='solid'
        ) 