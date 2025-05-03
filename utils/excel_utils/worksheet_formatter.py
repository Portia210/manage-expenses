"""Worksheet formatting utilities."""
from typing import List
import pandas as pd
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from .style_config import StyleFactory, ColumnConfig

class WorksheetFormatter:
    """Handles formatting of Excel worksheets."""
    
    def __init__(self, worksheet: Worksheet):
        self.worksheet = worksheet
        self.header_style = StyleFactory.create_header_fill()
        self.first_col_style = StyleFactory.create_first_column_fill()
    
    def apply_header_style(self, start_row: int, start_col: int, num_cols: int) -> None:
        """Apply header styling to the specified range."""
        for col_num in range(num_cols):
            cell = self.worksheet.cell(
                row=start_row + 1,
                column=start_col + col_num + 1
            )
            cell.fill = self.header_style
    
    def apply_first_column_style(self, start_row: int, start_col: int, num_rows: int) -> None:
        """Apply first column styling to the specified range."""
        for row_num in range(num_rows):
            cell = self.worksheet.cell(
                row=start_row + row_num + 2,
                column=start_col + 1
            )
            cell.fill = self.first_col_style
    
    def adjust_column_widths(self, df: pd.DataFrame, start_col: int) -> None:
        """Adjust column widths based on content."""
        for col_num in range(len(df.columns)):
            col_letter = openpyxl.utils.get_column_letter(start_col + col_num + 1)
            max_length = max(
                df[df.columns[col_num]].astype(str).apply(len).max(),
                len(str(df.columns[col_num]))
            )
            adjusted_width = min(
                max_length + ColumnConfig.DEFAULT_PADDING,
                ColumnConfig.MAX_WIDTH
            )
            self.worksheet.column_dimensions[col_letter].width = adjusted_width
    
    def format_dataframe(self, df: pd.DataFrame, start_row: int, start_col: int) -> None:
        """Apply all formatting to a DataFrame region."""
        self.apply_header_style(start_row, start_col, len(df.columns))
        self.apply_first_column_style(start_row, start_col, len(df))
        self.adjust_column_widths(df, start_col) 